from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from db import query, execute
from config import Config
from datetime import date
from datetime import datetime, timedelta
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename
import shutil, random

os.environ['TZ'] = 'Asia/Jakarta'

# Konfigurasi upload
UPLOAD_FOLDER = os.path.join('static', 'uploads', 'soal')
ALLOWED_IMG = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
ALLOWED_XLSX = {'xlsx'}

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

SOAL_UPLOAD_ROOT = os.path.join('static', 'uploads', 'soal')
os.makedirs(SOAL_UPLOAD_ROOT, exist_ok=True)


def generate_folder_code():
    """Generate kode unik 10 karakter untuk folder soal: S + 9 hex"""
    while True:
        kode = 'S' + uuid.uuid4().hex[:9].upper()
        # Pastikan tidak bentrok di DB
        cek = query("SELECT id FROM soal WHERE folder=%s", (kode,), one=True)
        if not cek:
            return kode


def get_soal_folder(folder_code, create=False):
    """Return path folder soal, buat jika create=True"""
    if not folder_code:
        return None
    path = os.path.join(SOAL_UPLOAD_ROOT, folder_code)
    if create and not os.path.exists(path):
        os.makedirs(path, exist_ok=True)
    return path


def hapus_folder_soal(folder_code):
    """Hapus folder + semua isinya"""
    if not folder_code:
        return
    path = os.path.join(SOAL_UPLOAD_ROOT, folder_code)
    if os.path.exists(path):
        shutil.rmtree(path)


def resolve_gambar(folder, filename):
    """Return relative URL path untuk template: 'uploads/soal/{folder}/{file}'"""
    if not folder or not filename:
        return None
    return f"uploads/soal/{folder}/{filename}"


def list_gambar_soal(folder):
    """List semua file gambar di folder soal"""
    if not folder:
        return []
    path = os.path.join(SOAL_UPLOAD_ROOT, folder)
    if not os.path.exists(path):
        return []
    return sorted([f for f in os.listdir(path)
                   if os.path.isfile(os.path.join(path, f))
                   and f.rsplit('.', 1)[-1].lower() in ALLOWED_IMG])
                   
def allowed_file(filename, allowed):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed

app = Flask(__name__)
app.config.from_object(Config)

# ============ DECORATORS ============
def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if 'user_id' not in session:
            flash('Silakan login terlebih dahulu', 'error')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return wrapper

def role_required(*roles):
    def decorator(f):
        @wraps(f)
        def wrapper(*args, **kwargs):
            if 'user_id' not in session:
                return redirect(url_for('login'))
            if session.get('role') not in roles:
                flash('Anda tidak punya akses ke halaman ini', 'error')
                return redirect(url_for('dashboard'))
            return f(*args, **kwargs)
        return wrapper
    return decorator

@app.context_processor
def inject_user():
    return dict(
        current_user=session.get('username'),
        current_role=session.get('role'),
        active_page=request.endpoint,
    )

# ============ LANDING ============
@app.route('/')
def landing():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))

    # Query statistik real dari database
    stats = {
        'siswa': query("SELECT COUNT(*) AS c FROM siswa", one=True)['c'],
        'guru':  query("SELECT COUNT(*) AS c FROM guru",  one=True)['c'],
        'kelas': query("SELECT COUNT(*) AS c FROM kelas", one=True)['c'],
        'soal':  query("SELECT COUNT(*) AS c FROM soal",  one=True)['c'],
        'ujian': query("SELECT COUNT(*) AS c FROM ujian WHERE status='aktif'", one=True)['c'],
    }

    return render_template('landing.html', stats=stats)

# ============ AUTH ============
@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form['username'].strip()
        password = request.form['password']
        user = query("SELECT * FROM users WHERE username=%s", (username,), one=True)
        if user and check_password_hash(user['password'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['role'] = user['role']
            flash(f"Selamat datang, {username}!", 'success')
            return redirect(url_for('dashboard'))
        flash('Username atau password salah!', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    flash('Anda telah logout', 'success')
    return redirect(url_for('login'))

# ============ DASHBOARD ============
@app.route('/dashboard')
@login_required
def dashboard():
    stats = {
        'siswa': query("SELECT COUNT(*) c FROM siswa", one=True)['c'],
        'guru': query("SELECT COUNT(*) c FROM guru", one=True)['c'],
        'kelas': query("SELECT COUNT(*) c FROM kelas", one=True)['c'],
        'soal': query("SELECT COUNT(*) c FROM soal", one=True)['c'],
        'ujian': query("SELECT COUNT(*) c FROM ujian", one=True)['c'],
    }
    ujian_terbaru = query("""
        SELECT u.*, k.nama_kelas FROM ujian u
        LEFT JOIN kelas k ON u.kelas_id=k.id
        ORDER BY u.id DESC LIMIT 5
    """)
    return render_template('dashboard.html', stats=stats, ujian_terbaru=ujian_terbaru)

# ============ KELAS ============
@app.route('/kelas', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def kelas():
    if request.method == 'POST':
        id_ = request.form.get('id')
        nama_kelas = request.form['nama_kelas']
        tingkat = request.form['tingkat']
        if id_:
            execute("UPDATE kelas SET nama_kelas=%s, tingkat=%s WHERE id=%s",
                    (nama_kelas, tingkat, id_))
            flash('Kelas berhasil diupdate', 'success')
        else:
            execute("INSERT INTO kelas (nama_kelas, tingkat) VALUES (%s,%s)",
                    (nama_kelas, tingkat))
            flash('Kelas berhasil ditambahkan', 'success')
        return redirect(url_for('kelas'))

    edit_data = None
    if request.args.get('edit'):
        edit_data = query("SELECT * FROM kelas WHERE id=%s",
                          (request.args['edit'],), one=True)
    data = query("SELECT * FROM kelas ORDER BY id DESC")
    return render_template('kelas.html', data=data, edit=edit_data)

@app.route('/kelas/hapus/<int:id>')
@login_required
@role_required('admin')
def kelas_hapus(id):
    execute("DELETE FROM kelas WHERE id=%s", (id,))
    flash('Kelas dihapus', 'success')
    return redirect(url_for('kelas'))

# ============ SISWA ============
@app.route('/siswa', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def siswa():
    if request.method == 'POST':
        id_ = request.form.get('id')
        nis = request.form['nis']
        nama = request.form['nama']
        kelas_id = request.form['kelas_id']
        jk = request.form['jk']

        if id_:
            execute("UPDATE siswa SET nama=%s, kelas_id=%s, jk=%s WHERE id=%s",
                    (nama, kelas_id, jk, id_))
            flash('Data siswa diupdate', 'success')
        else:
            # auto buat user
            uid, _ = execute(
                "INSERT INTO users (username, password, role) VALUES (%s,%s,'siswa')",
                (nis, generate_password_hash(nis))
            )
            execute("INSERT INTO siswa (user_id, nis, nama, kelas_id, jk) VALUES (%s,%s,%s,%s,%s)",
                    (uid, nis, nama, kelas_id, jk))
            flash('Siswa ditambahkan (login: NIS/NIS)', 'success')
        return redirect(url_for('siswa'))

    edit_data = None
    if request.args.get('edit'):
        edit_data = query("SELECT * FROM siswa WHERE id=%s",
                          (request.args['edit'],), one=True)
    data = query("""
        SELECT s.*, k.nama_kelas FROM siswa s
        LEFT JOIN kelas k ON s.kelas_id=k.id
        ORDER BY s.id DESC
    """)
    kelas_list = query("SELECT * FROM kelas ORDER BY nama_kelas")
    return render_template('siswa.html', data=data, kelas_list=kelas_list, edit=edit_data)

@app.route('/siswa/hapus/<int:id>')
@login_required
@role_required('admin')
def siswa_hapus(id):
    s = query("SELECT user_id FROM siswa WHERE id=%s", (id,), one=True)
    if s and s['user_id']:
        execute("DELETE FROM users WHERE id=%s", (s['user_id'],))
    execute("DELETE FROM siswa WHERE id=%s", (id,))
    flash('Siswa dihapus', 'success')
    return redirect(url_for('siswa'))

# ============ GURU ============
@app.route('/guru', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def guru():
    if request.method == 'POST':
        id_ = request.form.get('id')
        nip = request.form['nip']
        nama = request.form['nama']
        mapel = request.form['mapel']
        jk = request.form['jk']

        if id_:
            execute("UPDATE guru SET nama=%s, mapel=%s, jk=%s WHERE id=%s",
                    (nama, mapel, jk, id_))
            flash('Data guru diupdate', 'success')
        else:
            uid, _ = execute(
                "INSERT INTO users (username, password, role) VALUES (%s,%s,'guru')",
                (nip, generate_password_hash(nip))
            )
            execute("INSERT INTO guru (user_id, nip, nama, mapel, jk) VALUES (%s,%s,%s,%s,%s)",
                    (uid, nip, nama, mapel, jk))
            flash('Guru ditambahkan (login: NIP/NIP)', 'success')
        return redirect(url_for('guru'))

    edit_data = None
    if request.args.get('edit'):
        edit_data = query("SELECT * FROM guru WHERE id=%s",
                          (request.args['edit'],), one=True)
    data = query("SELECT * FROM guru ORDER BY id DESC")
    return render_template('guru.html', data=data, edit=edit_data)

@app.route('/guru/hapus/<int:id>')
@login_required
@role_required('admin')
def guru_hapus(id):
    g = query("SELECT user_id FROM guru WHERE id=%s", (id,), one=True)
    if g and g['user_id']:
        execute("DELETE FROM users WHERE id=%s", (g['user_id'],))
    execute("DELETE FROM guru WHERE id=%s", (id,))
    flash('Guru dihapus', 'success')
    return redirect(url_for('guru'))

# ============ SOAL ============
@app.route('/soal', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'guru')
def soal():
    if request.method == 'POST':
        id_ = request.form.get('id')
        pertanyaan = request.form['pertanyaan']
        a, b, c, d = (request.form[f'opsi_{x}'] for x in 'abcd')
        jawaban = request.form['jawaban']

        if id_:
            # Update — folder tidak berubah
            execute("""UPDATE soal SET pertanyaan=%s, opsi_a=%s, opsi_b=%s,
                       opsi_c=%s, opsi_d=%s, jawaban=%s WHERE id=%s""",
                    (pertanyaan, a, b, c, d, jawaban, id_))
            flash('Soal diupdate.', 'success')
            return redirect(url_for('soal_gambar', soal_id=id_))
        else:
            # Insert baru → generate folder unik
            folder = generate_folder_code()
            soal_id, _ = execute("""INSERT INTO soal
                (folder, pertanyaan, opsi_a, opsi_b, opsi_c, opsi_d, jawaban, sumber)
                VALUES (%s,%s,%s,%s,%s,%s,%s,'manual')""",
                (folder, pertanyaan, a, b, c, d, jawaban))
            # Buat folder fisik
            get_soal_folder(folder, create=True)
            flash('Soal berhasil dibuat. Silakan upload gambarnya.', 'success')
            return redirect(url_for('soal_gambar', soal_id=soal_id))

    edit_data = None
    if request.args.get('edit'):
        edit_data = query("SELECT * FROM soal WHERE id=%s",
                          (request.args['edit'],), one=True)
    data = query("SELECT * FROM soal ORDER BY id DESC")
    return render_template('soal.html', data=data, edit=edit_data)
    
@app.route('/soal/hapus/<int:id>')
@login_required
@role_required('admin', 'guru')
def soal_hapus(id):
    s = query("SELECT folder FROM soal WHERE id=%s", (id,), one=True)
    if s:
        hapus_folder_soal(s['folder'])
    execute("DELETE FROM soal WHERE id=%s", (id,))
    flash('Soal & folder gambarnya dihapus.', 'success')
    return redirect(url_for('soal'))

# ============ UJIAN ============
@app.route('/ujian', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'guru')
def ujian():
    if request.method == 'POST':
        id_ = request.form.get('id')
        judul = request.form['judul']
        mapel = request.form['mapel']
        kelas_id = request.form['kelas_id']
        durasi = request.form['durasi']
        tanggal = request.form['tanggal']
        status = request.form['status']
        soal_ids = request.form.getlist('soal_ids')
        shuffle_soal = 1 if request.form.get('shuffle_soal') else 0
        shuffle_opsi = 1 if request.form.get('shuffle_opsi') else 0
        
        if id_:
            execute("""UPDATE ujian SET judul=%s, mapel=%s, kelas_id=%s, durasi=%s,
                       tanggal=%s, status=%s, shuffle_soal=%s, shuffle_opsi=%s WHERE id=%s""",
                (judul, mapel, kelas_id, durasi, tanggal, status,
                 shuffle_soal, shuffle_opsi, id_))
            ujian_id = id_
            flash('Ujian diupdate', 'success')
        else:
            ujian_id, _ = execute("""INSERT INTO ujian
                    (judul, mapel, kelas_id, durasi, tanggal, status, shuffle_soal, shuffle_opsi)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (judul, mapel, kelas_id, durasi, tanggal, status, shuffle_soal, shuffle_opsi))
            flash('Ujian dibuat', 'success')

        for sid in soal_ids:
            execute("INSERT INTO ujian_soal (ujian_id, soal_id) VALUES (%s,%s)",
                    (ujian_id, sid))
        return redirect(url_for('ujian'))

    edit_data = None
    selected_soal = []
    if request.args.get('edit'):
        edit_data = query("SELECT * FROM ujian WHERE id=%s",
                          (request.args['edit'],), one=True)
        selected_soal = [r['soal_id'] for r in query(
            "SELECT soal_id FROM ujian_soal WHERE ujian_id=%s",
            (request.args['edit'],))]

    data = query("""
        SELECT u.*, k.nama_kelas,
            (SELECT COUNT(*) FROM ujian_soal WHERE ujian_id=u.id) AS jml_soal
        FROM ujian u
        LEFT JOIN kelas k ON u.kelas_id=k.id
        ORDER BY u.id DESC
    """)
    kelas_list = query("SELECT * FROM kelas ORDER BY nama_kelas")
    soal_list = query("SELECT * FROM soal ORDER BY id DESC")
    return render_template('ujian.html', data=data, kelas_list=kelas_list,
                           soal_list=soal_list, edit=edit_data,
                           selected_soal=selected_soal)

@app.route('/ujian/hapus/<int:id>')
@login_required
@role_required('admin', 'guru')
def ujian_hapus(id):
    execute("DELETE FROM ujian WHERE id=%s", (id,))
    flash('Ujian dihapus', 'success')
    return redirect(url_for('ujian'))

# ============ CBT (Siswa) ============
@app.route('/cbt')
@login_required
def cbt_list():
    data = query("""
        SELECT u.*, k.nama_kelas FROM ujian u
        LEFT JOIN kelas k ON u.kelas_id=k.id
        WHERE u.status='aktif' ORDER BY u.id DESC
    """)
    return render_template('cbt_list.html', data=data)

@app.route('/cbt/<int:ujian_id>', methods=['GET', 'POST'])
@login_required
def cbt_form(ujian_id):
    ujian = query("SELECT * FROM ujian WHERE id=%s", (ujian_id,), one=True)
    if not ujian:
        flash('Ujian tidak ditemukan', 'error')
        return redirect(url_for('cbt_list'))

    soal_list = query("""
        SELECT s.* FROM soal s
        JOIN ujian_soal us ON s.id=us.soal_id
        WHERE us.ujian_id=%s
    """, (ujian_id,))

    if request.method == 'POST':
        jawaban = request.form  # format: jawaban_<soal_id>
        benar = salah = 0
        for s in soal_list:
            jw = jawaban.get(f'jawaban_{s["id"]}', '')
            if jw == s['jawaban']:
                benar += 1
            else:
                salah += 1
        total = len(soal_list)
        nilai = round(benar / total * 100) if total else 0

        # ambil siswa_id
        siswa = query("SELECT id FROM siswa WHERE user_id=%s",
                      (session['user_id'],), one=True)
        siswa_id = siswa['id'] if siswa else None

        hasil_id, _ = execute("""INSERT INTO hasil_ujian
            (ujian_id, siswa_id, nilai, benar, salah) VALUES (%s,%s,%s,%s,%s)""",
            (ujian_id, siswa_id, nilai, benar, salah))

        for s in soal_list:
            execute("INSERT INTO jawaban_siswa (hasil_id, soal_id, jawaban) VALUES (%s,%s,%s)",
                    (hasil_id, s['id'], jawaban.get(f'jawaban_{s["id"]}', '')))

        return redirect(url_for('cbt_hasil', hasil_id=hasil_id))

    return render_template('cbt_form.html', ujian=ujian, soal_list=soal_list)

@app.route('/cbt/hasil/<int:hasil_id>')
@login_required
def cbt_hasil(hasil_id):
    h = query("""SELECT h.*, u.judul FROM hasil_ujian h
                 JOIN ujian u ON h.ujian_id=u.id
                 WHERE h.id=%s""", (hasil_id,), one=True)
    return render_template('cbt_hasil.html', hasil=h)

# ============ API UNTUK TIMER ============
@app.route('/api/durasi/<int:ujian_id>')
@login_required
def api_durasi(ujian_id):
    u = query("SELECT durasi FROM ujian WHERE id=%s", (ujian_id,), one=True)
    return jsonify({'durasi': u['durasi'] * 60 if u else 0})


# ============ STUDENT LOGIN GROUP ============
@app.route('/login/siswa', methods=['GET', 'POST'])
def login_siswa():
    if 'user_id' in session and session.get('role') == 'siswa':
        return redirect(url_for('siswa_dashboard'))
    if request.method == 'POST':
        username = request.form['username'].strip()
        password = request.form['password']
        user = query("SELECT * FROM users WHERE username=%s AND role='siswa'",
                     (username,), one=True)
        if user and check_password_hash(user['password'], password):
            # Ambil data siswa
            siswa = query("SELECT s.*, k.nama_kelas FROM siswa s "
                          "LEFT JOIN kelas k ON s.kelas_id=k.id "
                          "WHERE s.user_id=%s", (user['id'],), one=True)
            if not siswa:
                flash('Data siswa tidak ditemukan. Hubungi admin.', 'error')
                return render_template('siswa_login.html')
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['role'] = 'siswa'
            session['siswa_id'] = siswa['id']
            session['siswa_nama'] = siswa['nama']
            session['siswa_kelas'] = siswa['nama_kelas']
            session['kelas_id'] = siswa['kelas_id']
            flash(f"Selamat datang, {siswa['nama']}!", 'success')
            return redirect(url_for('siswa_dashboard'))
        flash('NIS atau password salah!', 'error')
    return render_template('siswa_login.html')


# ============ STUDENT DASHBOARD ============
@app.route('/siswa/dashboard')
@login_required
def siswa_dashboard():
    if session.get('role') != 'siswa':
        flash('Akses khusus siswa.', 'error')
        return redirect(url_for('dashboard'))

    siswa_id = session.get('siswa_id')
    kelas_id = session.get('kelas_id')

    # Daftar ujian untuk kelas siswa
    ujian_list = query("""
        SELECT u.*, k.nama_kelas,
            (SELECT COUNT(*) FROM ujian_soal WHERE ujian_id=u.id) AS jml_soal,
            (SELECT id FROM hasil_ujian WHERE ujian_id=u.id AND siswa_id=%s 
                AND status='selesai' ORDER BY id DESC LIMIT 1) AS hasil_selesai_id,
            (SELECT id FROM hasil_ujian WHERE ujian_id=u.id AND siswa_id=%s 
                AND status='berlangsung' ORDER BY id DESC LIMIT 1) AS hasil_berlangsung_id
        FROM ujian u
        LEFT JOIN kelas k ON u.kelas_id=k.id
        WHERE u.status='aktif' AND u.kelas_id=%s
        ORDER BY u.tanggal DESC, u.id DESC
    """, (siswa_id, siswa_id, kelas_id))

    # Statistik
    stats = {
        'total_ujian': len(ujian_list),
        'selesai': sum(1 for u in ujian_list if u['hasil_selesai_id']),
        'tersedia': sum(1 for u in ujian_list 
                        if not u['hasil_selesai_id'] and not u['hasil_berlangsung_id']),
        'berlangsung': sum(1 for u in ujian_list if u['hasil_berlangsung_id']),
    }

    return render_template('siswa_dashboard.html', 
                           ujian_list=ujian_list, stats=stats)


# ============ START / MULAI UJIAN ============
@app.route('/siswa/ujian/<int:ujian_id>/start')
@login_required
def siswa_ujian_start(ujian_id):
    if session.get('role') != 'siswa':
        return redirect(url_for('dashboard'))

    siswa_id = session.get('siswa_id')
    ujian = query("SELECT * FROM ujian WHERE id=%s AND status='aktif'",
                  (ujian_id,), one=True)
    if not ujian:
        flash('Ujian tidak ditemukan atau tidak aktif.', 'error')
        return redirect(url_for('siswa_dashboard'))

    # Cek sudah selesai
    sudah = query("SELECT id FROM hasil_ujian WHERE ujian_id=%s AND siswa_id=%s "
                  "AND status='selesai' LIMIT 1",
                  (ujian_id, siswa_id), one=True)
    if sudah:
        flash('Anda sudah menyelesaikan ujian ini.', 'error')
        return redirect(url_for('siswa_hasil', hasil_id=sudah['id']))

    # Cek sesi berlangsung
    berlangsung = query("SELECT * FROM hasil_ujian WHERE ujian_id=%s AND siswa_id=%s "
                        "AND status='berlangsung' ORDER BY id DESC LIMIT 1",
                        (ujian_id, siswa_id), one=True)

    if berlangsung:
        from datetime import datetime as dt
        if dt.now() >= berlangsung['waktu_expired']:
            return redirect(url_for('siswa_ujian_submit',
                                    ujian_id=ujian_id, auto='1'))
        return redirect(url_for('siswa_ujian', ujian_id=ujian_id))

    # === BUAT SESI BARU + RANDOMISASI ===
    from datetime import datetime, timedelta
    waktu_mulai = datetime.now()
    waktu_expired = waktu_mulai + timedelta(minutes=ujian['durasi'])

    # Ambil semua soal ujian
    daftar_soal = query("""
        SELECT soal_id FROM ujian_soal WHERE ujian_id=%s
    """, (ujian_id,))
    soal_ids = [s['soal_id'] for s in daftar_soal]

    if not soal_ids:
        flash('Ujian belum memiliki soal.', 'error')
        return redirect(url_for('siswa_dashboard'))

    # Shuffle soal
    if ujian.get('shuffle_soal', 1):
        random.shuffle(soal_ids)

    # Buat record hasil_ujian
    hasil_id, _ = execute("""
        INSERT INTO hasil_ujian (ujian_id, siswa_id, nilai, benar, salah,
                                 waktu_mulai, waktu_expired, status)
        VALUES (%s, %s, 0, 0, 0, %s, %s, 'berlangsung')
    """, (ujian_id, siswa_id, waktu_mulai, waktu_expired))

    # Simpan urutan soal + urutan opsi (acak)
    for urutan, sid in enumerate(soal_ids, start=1):
        if ujian.get('shuffle_opsi', 1):
            opsi = list('ABCD')
            random.shuffle(opsi)
            urutan_opsi = ''.join(opsi)
        else:
            urutan_opsi = 'ABCD'

        execute("""INSERT INTO ujian_siswa_soal
                   (hasil_id, soal_id, urutan, urutan_opsi)
                   VALUES (%s, %s, %s, %s)""",
                (hasil_id, sid, urutan, urutan_opsi))

    return redirect(url_for('siswa_ujian', ujian_id=ujian_id))

@app.route('/siswa/ujian/<int:ujian_id>', methods=['GET', 'POST'])
@login_required
def siswa_ujian(ujian_id):
    if session.get('role') != 'siswa':
        return redirect(url_for('dashboard'))

    from datetime import datetime
    siswa_id = session.get('siswa_id')
    ujian = query("SELECT * FROM ujian WHERE id=%s", (ujian_id,), one=True)

    sesi = query("""SELECT * FROM hasil_ujian
                    WHERE ujian_id=%s AND siswa_id=%s AND status='berlangsung'
                    ORDER BY id DESC LIMIT 1""",
                 (ujian_id, siswa_id), one=True)

    if not sesi:
        flash('Sesi ujian tidak ditemukan.', 'error')
        return redirect(url_for('siswa_dashboard'))

    if datetime.now() >= sesi['waktu_expired']:
        return redirect(url_for('siswa_ujian_submit',
                                ujian_id=ujian_id, auto='1'))

    # Ambil soal SESUAI URUTAN YANG DISIMPAN
    soal_list = query("""
        SELECT s.*, uss.urutan, uss.urutan_opsi
        FROM ujian_siswa_soal uss
        JOIN soal s ON s.id = uss.soal_id
        WHERE uss.hasil_id=%s
        ORDER BY uss.urutan ASC
    """, (sesi['id'],))

    jawaban_tersimpan = query("""
        SELECT soal_id, jawaban FROM jawaban_siswa WHERE hasil_id=%s
    """, (sesi['id'],))
    jawaban_map = {j['soal_id']: j['jawaban'] for j in jawaban_tersimpan}

    if request.method == 'POST':
        return redirect(url_for('siswa_ujian_submit', ujian_id=ujian_id))

    sisa_detik = int((sesi['waktu_expired'] - datetime.now()).total_seconds())
    if sisa_detik < 0:
        sisa_detik = 0

    return render_template('siswa_ujian.html',
                           ujian=ujian, soal_list=soal_list,
                           sesi=sesi, sisa_detik=sisa_detik,
                           jawaban_map=jawaban_map)

# ============ SIMPAN JAWABAN SEMENTARA (AJAX) ============
@app.route('/siswa/ujian/<int:ujian_id>/simpan', methods=['POST'])
@login_required
def siswa_ujian_simpan(ujian_id):
    if session.get('role') != 'siswa':
        return jsonify({'ok': False, 'msg': 'Akses ditolak'}), 403

    siswa_id = session.get('siswa_id')
    sesi = query("""SELECT * FROM hasil_ujian 
                    WHERE ujian_id=%s AND siswa_id=%s AND status='berlangsung'
                    ORDER BY id DESC LIMIT 1""",
                 (ujian_id, siswa_id), one=True)
    if not sesi:
        return jsonify({'ok': False, 'msg': 'Sesi tidak ditemukan'}), 404

    data = request.get_json() or {}
    soal_id = data.get('soal_id')
    jawaban = data.get('jawaban')  # 'A','B','C','D'

    # Hapus jawaban lama untuk soal ini, lalu insert
    execute("DELETE FROM jawaban_siswa WHERE hasil_id=%s AND soal_id=%s",
            (sesi['id'], soal_id))
    if jawaban:
        execute("INSERT INTO jawaban_siswa (hasil_id, soal_id, jawaban) VALUES (%s,%s,%s)",
                (sesi['id'], soal_id, jawaban))

    return jsonify({'ok': True})


@app.route('/siswa/ujian/<int:ujian_id>/submit')
@login_required
def siswa_ujian_submit(ujian_id):
    if session.get('role') != 'siswa':
        return redirect(url_for('dashboard'))

    siswa_id = session.get('siswa_id')
    sesi = query("""SELECT * FROM hasil_ujian
                    WHERE ujian_id=%s AND siswa_id=%s AND status='berlangsung'
                    ORDER BY id DESC LIMIT 1""",
                 (ujian_id, siswa_id), one=True)
    if not sesi:
        flash('Sesi ujian tidak ditemukan.', 'error')
        return redirect(url_for('siswa_dashboard'))

    # Ambil soal + urutan opsi dari sesi
    soal_list = query("""
        SELECT s.id, s.jawaban AS jawaban_asli, uss.urutan_opsi
        FROM ujian_siswa_soal uss
        JOIN soal s ON s.id = uss.soal_id
        WHERE uss.hasil_id=%s
    """, (sesi['id'],))

    jawaban_list = query("""
        SELECT soal_id, jawaban FROM jawaban_siswa WHERE hasil_id=%s
    """, (sesi['id'],))
    jawaban_map = {j['soal_id']: j['jawaban'] for j in jawaban_list}

    # Konversi: urutan_opsi='CADB' artinya:
    #   label 'A' → opsi asli 'C'
    #   label 'B' → opsi asli 'A'
    #   label 'C' → opsi asli 'D'
    #   label 'D' → opsi asli 'B'
    LABELS = ['A', 'B', 'C', 'D']

    benar = salah = 0
    for s in soal_list:
        urutan_opsi = s['urutan_opsi'] or 'ABCD'
        jw_label = jawaban_map.get(s['id'], '')  # label yang dipilih siswa
        if jw_label in LABELS:
            idx = LABELS.index(jw_label)
            jw_asli = urutan_opsi[idx]
        else:
            jw_asli = ''

        if jw_asli and jw_asli == s['jawaban_asli']:
            benar += 1
        else:
            salah += 1

    total = len(soal_list)
    nilai = round(benar / total * 100) if total else 0

    execute("""UPDATE hasil_ujian
               SET nilai=%s, benar=%s, salah=%s, status='selesai',
                   waktu_selesai=NOW()
               WHERE id=%s""",
            (nilai, benar, salah, sesi['id']))

    return redirect(url_for('siswa_hasil', hasil_id=sesi['id']))

# ============ HASIL UJIAN SISWA ============
@app.route('/siswa/hasil/<int:hasil_id>')
@login_required
def siswa_hasil(hasil_id):
    if session.get('role') != 'siswa':
        return redirect(url_for('dashboard'))

    siswa_id = session.get('siswa_id')
    h = query("""SELECT h.*, u.judul, u.mapel, u.durasi, k.nama_kelas
                 FROM hasil_ujian h
                 JOIN ujian u ON h.ujian_id=u.id
                 LEFT JOIN kelas k ON u.kelas_id=k.id
                 WHERE h.id=%s AND h.siswa_id=%s""",
              (hasil_id, siswa_id), one=True)
    if not h:
        flash('Hasil tidak ditemukan.', 'error')
        return redirect(url_for('siswa_dashboard'))

    # Detail dengan urutan opsi
    detail_raw = query("""
        SELECT s.*, uss.urutan_opsi, uss.urutan,
               COALESCE(js.jawaban,'') AS jawaban_siswa
        FROM ujian_siswa_soal uss
        JOIN soal s ON s.id = uss.soal_id
        LEFT JOIN jawaban_siswa js ON js.soal_id=s.id AND js.hasil_id=%s
        WHERE uss.hasil_id=%s
        ORDER BY uss.urutan ASC
    """, (hasil_id, hasil_id))

    LABELS = ['A', 'B', 'C', 'D']
    detail = []
    for d in detail_raw:
        urutan = d['urutan_opsi'] or 'ABCD'
        opsi_display = []
        for i, label in enumerate(LABELS):
            key_asli = urutan[i].lower()
            gmb_file = d.get(f'gambar_{key_asli}')
            opsi_display.append({
                'label': label,
                'teks': d.get(f'opsi_{key_asli}', '') or '',
                'gambar': gmb_file,
                'gambar_url': resolve_gambar(d['folder'], gmb_file),
                'adalah_jawaban': urutan[i] == d['jawaban'],
                'dipilih': label == d['jawaban_siswa'],
            })
        detail.append({
            'pertanyaan': d['pertanyaan'],
            'gambar': d.get('gambar'),
            'gambar_url': resolve_gambar(d['folder'], d.get('gambar')),
            'opsi': opsi_display,
        })

    return render_template('siswa_hasil.html', hasil=h, detail=detail)

@app.route('/soal/template')
@login_required
@role_required('admin', 'guru')
def soal_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Soal"

    # Header styling
    headers = ['pertanyaan', 'gambar',
               'opsi_a', 'gambar_a',
               'opsi_b', 'gambar_b',
               'opsi_c', 'gambar_c',
               'opsi_d', 'gambar_d',
               'jawaban']
    header_fill = PatternFill("solid", fgColor="16A34A")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 22

    # Baris contoh
    contoh = [
        'Ibu kota Indonesia adalah?', '',
        'Jakarta', '', 'Bandung', '', 'Surabaya', '', 'Medan', '',
        'A'
    ]
    for col, val in enumerate(contoh, start=1):
        ws.cell(row=2, column=col, value=val)

    # Baris contoh dengan gambar
    contoh2 = [
        'Perhatikan gambar berikut. Bangun datar apa ini?', 'segitiga.png',
        'Segitiga', '', 'Persegi', '', 'Lingkaran', '', 'Trapesium', '',
        'A'
    ]
    for col, val in enumerate(contoh2, start=1):
        ws.cell(row=3, column=col, value=val)

    # Instruksi
    ws2 = wb.create_sheet("PETUNJUK")
    petunjuk = [
    ["PETUNJUK PENGISIAN TEMPLATE SOAL"],
    [""],
    ["1. Sheet 'Bank Soal' digunakan untuk mengisi soal."],
    ["2. Kolom WAJIB: pertanyaan, opsi_a, opsi_b, opsi_c, opsi_d, jawaban"],
    ["3. Kolom OPSIONAL: gambar, gambar_a, gambar_b, gambar_c, gambar_d"],
    ["4. Isi kolom gambar dengan NAMA FILE saja (mis. segitiga.png)."],
    ["   JANGAN isi path/folder — sistem akan otomatis menyimpannya"],
    ["   ke folder unik yang dibuat per soal."],
    ["5. Format jawaban: A / B / C / D (huruf kapital)."],
    ["6. Jangan ubah nama header di baris pertama."],
    ["7. Baris 2 & 3 adalah contoh — silakan timpa/hapus."],
    [""],
    ["ALUR:"],
    ["  a. Upload file Excel ini → sistem buat folder unik per soal"],
    ["  b. Setelah import, Anda otomatis diarahkan ke halaman upload gambar"],
    ["  c. Upload gambar-gambar yang namanya disebut di kolom gambar"],
    ["  d. Gambar otomatis masuk ke folder masing-masing soal"],
]
    for i, row in enumerate(petunjuk, start=1):
        ws2.cell(row=i, column=1, value=row[0]).font = Font(
            bold=(i == 1), size=12 if i == 1 else 11)
    ws2.column_dimensions['A'].width = 80

    from flask import send_file
    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name='template_soal_cbt.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/soal/upload', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'guru')
def soal_upload():
    if request.method == 'POST':
        file = request.files.get('file_excel')
        if not file or file.filename == '':
            flash('Pilih file Excel dulu.', 'error')
            return redirect(url_for('soal_upload'))
        if not allowed_file(file.filename, ALLOWED_XLSX):
            flash('File harus .xlsx', 'error')
            return redirect(url_for('soal_upload'))

        try:
            wb = load_workbook(file, data_only=True)
        except Exception as e:
            flash(f'Gagal baca Excel: {e}', 'error')
            return redirect(url_for('soal_upload'))

        if 'Bank Soal' not in wb.sheetnames:
            flash("Sheet 'Bank Soal' tidak ditemukan.", 'error')
            return redirect(url_for('soal_upload'))

        ws = wb['Bank Soal']
        expected = ['pertanyaan', 'gambar',
                    'opsi_a', 'gambar_a', 'opsi_b', 'gambar_b',
                    'opsi_c', 'gambar_c', 'opsi_d', 'gambar_d', 'jawaban']
        header_row = [c.value.strip().lower() if c.value else '' for c in ws[1]]
        col_idx = {name: (header_row.index(name) if name in header_row else None)
                   for name in expected}

        required = ['pertanyaan', 'opsi_a', 'opsi_b', 'opsi_c', 'opsi_d', 'jawaban']
        missing = [r for r in required if col_idx[r] is None]
        if missing:
            flash(f'Kolom wajib tidak ada: {", ".join(missing)}', 'error')
            return redirect(url_for('soal_upload'))

        def get_val(row, key):
            idx = col_idx.get(key)
            if idx is None:
                return ''
            v = row[idx].value if idx < len(row) else None
            return str(v).strip() if v is not None else ''

        sukses, gagal, inserted_ids = 0, [], []

        for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            pertanyaan = get_val(row, 'pertanyaan')
            if not pertanyaan:
                continue

            opsi_a = get_val(row, 'opsi_a')
            opsi_b = get_val(row, 'opsi_b')
            opsi_c = get_val(row, 'opsi_c')
            opsi_d = get_val(row, 'opsi_d')
            jawaban = get_val(row, 'jawaban').upper()

            if not all([opsi_a, opsi_b, opsi_c, opsi_d]):
                gagal.append(f'Baris {r_idx}: opsi tidak lengkap')
                continue
            if jawaban not in ['A', 'B', 'C', 'D']:
                gagal.append(f'Baris {r_idx}: jawaban harus A/B/C/D')
                continue

            # Generate folder UNIK per soal
            folder = generate_folder_code()

            gambar = get_val(row, 'gambar')
            gambar_a = get_val(row, 'gambar_a')
            gambar_b = get_val(row, 'gambar_b')
            gambar_c = get_val(row, 'gambar_c')
            gambar_d = get_val(row, 'gambar_d')

            soal_id, _ = execute("""INSERT INTO soal
                (folder, pertanyaan, gambar, opsi_a, gambar_a,
                 opsi_b, gambar_b, opsi_c, gambar_c,
                 opsi_d, gambar_d, jawaban, sumber)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'excel')""",
                (folder, pertanyaan, gambar or None,
                 opsi_a, gambar_a or None,
                 opsi_b, gambar_b or None,
                 opsi_c, gambar_c or None,
                 opsi_d, gambar_d or None,
                 jawaban))
            # Buat folder fisik
            get_soal_folder(folder, create=True)
            inserted_ids.append(soal_id)
            sukses += 1

        if sukses:
            flash(f'✅ {sukses} soal berhasil diimport. Silakan upload gambarnya.', 'success')
        if gagal:
            flash('⚠️ Gagal: ' + ' | '.join(gagal[:5]) +
                  (f' (+{len(gagal)-5} lagi)' if len(gagal) > 5 else ''), 'error')

        # Redirect ke halaman kelola gambar (batch) jika sukses
        if inserted_ids:
            return redirect(url_for('soal_gambar_batch', ids=','.join(map(str, inserted_ids))))
        return redirect(url_for('soal'))

    return render_template('soal_upload.html')

@app.route('/soal/<int:soal_id>/gambar', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'guru')
def soal_gambar(soal_id):
    soal = query("SELECT * FROM soal WHERE id=%s", (soal_id,), one=True)
    if not soal:
        flash('Soal tidak ditemukan.', 'error')
        return redirect(url_for('soal'))

    # Pastikan folder sudah ada (untuk soal lama)
    folder = soal['folder']
    if not folder:
        folder = generate_folder_code()
        execute("UPDATE soal SET folder=%s WHERE id=%s", (folder, soal_id))
        soal['folder'] = folder
    get_soal_folder(folder, create=True)

    if request.method == 'POST':
        files = request.files.getlist('gambar_files')
        sukses, gagal = 0, []
        for f in files:
            if not f or f.filename == '':
                continue
            if not allowed_file(f.filename, ALLOWED_IMG):
                gagal.append(f.filename)
                continue
            nama_aman = secure_filename(f.filename)
            target = os.path.join(get_soal_folder(folder), nama_aman)
            f.save(target)
            sukses += 1

        if sukses:
            flash(f'✅ {sukses} gambar berhasil diupload ke folder {folder}.', 'success')
        if gagal:
            flash(f'⚠️ Gagal: {", ".join(gagal)}', 'error')
        return redirect(url_for('soal_gambar', soal_id=soal_id))

    # List gambar di folder
    files = list_gambar_soal(folder)
    return render_template('soal_gambar.html', soal=soal, files=files)

@app.route('/soal/<int:soal_id>/gambar/hapus/<path:filename>')
@login_required
@role_required('admin', 'guru')
def soal_gambar_hapus(soal_id, filename):
    soal = query("SELECT folder FROM soal WHERE id=%s", (soal_id,), one=True)
    if not soal:
        flash('Soal tidak ditemukan.', 'error')
        return redirect(url_for('soal'))

    folder = soal['folder']
    nama_aman = secure_filename(filename)
    path = os.path.join(get_soal_folder(folder), nama_aman)
    if os.path.exists(path):
        os.remove(path)
        # Bersihkan referensi di kolom gambar
        execute("""UPDATE soal SET 
            gambar = IF(gambar=%s, NULL, gambar),
            gambar_a = IF(gambar_a=%s, NULL, gambar_a),
            gambar_b = IF(gambar_b=%s, NULL, gambar_b),
            gambar_c = IF(gambar_c=%s, NULL, gambar_c),
            gambar_d = IF(gambar_d=%s, NULL, gambar_d)
            WHERE id=%s""",
            (nama_aman, nama_aman, nama_aman, nama_aman, nama_aman, soal_id))
        flash(f'Gambar {nama_aman} dihapus.', 'success')
    return redirect(url_for('soal_gambar', soal_id=soal_id))

@app.route('/soal/gambar/batch')
@login_required
@role_required('admin', 'guru')
def soal_gambar_batch():
    ids_param = request.args.get('ids', '')
    if not ids_param:
        flash('Tidak ada soal yang dipilih.', 'error')
        return redirect(url_for('soal'))

    try:
        ids = [int(i) for i in ids_param.split(',') if i.strip().isdigit()]
    except ValueError:
        ids = []

    if not ids:
        flash('ID soal tidak valid.', 'error')
        return redirect(url_for('soal'))

    # Ambil semua soal yang baru diimport
    placeholder = ','.join(['%s'] * len(ids))
    soal_list = query(f"SELECT * FROM soal WHERE id IN ({placeholder}) ORDER BY id",
                      tuple(ids))

    # Info gambar per soal
    for s in soal_list:
        s['files'] = list_gambar_soal(s['folder'])

    return render_template('soal_gambar_batch.html', soal_list=soal_list)


@app.route('/soal/<int:soal_id>/gambar/quick', methods=['POST'])
@login_required
@role_required('admin', 'guru')
def soal_gambar_quick(soal_id):
    """Upload gambar cepat dari halaman batch (AJAX/redirect)"""
    soal = query("SELECT * FROM soal WHERE id=%s", (soal_id,), one=True)
    if not soal:
        return jsonify({'ok': False, 'msg': 'Soal tidak ditemukan'}), 404

    folder = soal['folder']
    get_soal_folder(folder, create=True)

    files = request.files.getlist('gambar_files')
    sukses = 0
    for f in files:
        if f and f.filename and allowed_file(f.filename, ALLOWED_IMG):
            nama = secure_filename(f.filename)
            f.save(os.path.join(get_soal_folder(folder), nama))
            sukses += 1

    flash(f'✅ {sukses} gambar diupload ke soal #{soal_id}.', 'success')
    return redirect(url_for('soal_gambar_batch',
                            ids=request.form.get('back_ids', '')))

@app.route('/soal/upload-gambar', methods=['POST'])
@login_required
@role_required('admin', 'guru')
def soal_upload_gambar():
    files = request.files.getlist('gambar_files')
    sukses, gagal = 0, []
    for f in files:
        if not f or f.filename == '':
            continue
        if not allowed_file(f.filename, ALLOWED_IMG):
            gagal.append(f.filename)
            continue
        # Amankan nama & cegah duplikat
        nama_aman = secure_filename(f.filename)
        base, ext = os.path.splitext(nama_aman)
        target = os.path.join(UPLOAD_FOLDER, nama_aman)
        i = 1
        while os.path.exists(target):
            nama_aman = f"{base}_{i}{ext}"
            target = os.path.join(UPLOAD_FOLDER, nama_aman)
            i += 1
        f.save(target)
        sukses += 1

    if sukses:
        flash(f'✅ {sukses} gambar berhasil diupload.', 'success')
    if gagal:
        flash(f'⚠️ Gagal upload: {", ".join(gagal)}', 'error')
    return redirect(url_for('soal_upload'))


# ============ LOGOUT SISWA ============
@app.route('/siswa/logout')
def siswa_logout():
    session.clear()
    flash('Anda telah logout.', 'success')
    return redirect(url_for('login_siswa'))
    
if __name__ == '__main__':
    app.run(debug=True, port=5005,host="0.0.0.0")